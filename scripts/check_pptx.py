from pptx import Presentation

p = Presentation('docs/ReclaimYourNet-KriticalHacks2026.pptx')
print('slides:', len(p.slides))
for i, s in enumerate(p.slides, 1):
    texts = [
        sh.text_frame.text.split('\n')[0][:56]
        for sh in s.shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    ]
    notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ''
    head = texts[0] if texts else '(title slide)'
    print(f'{i:2}. {head:58} | notes {len(notes):4} chars')
