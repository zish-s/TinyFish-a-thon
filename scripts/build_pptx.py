"""
Builds the submission .pptx from the deck in docs/PPT.md.

    python scripts/build_pptx.py

Produces docs/ReclaimYourNet-KriticalHacks2026.pptx - the file you attach to
the Unstop form. Slide notes carry the speaker script.
"""
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "PPT.md"
OUT = ROOT / "docs" / "ReclaimYourNet-KriticalHacks2026.pptx"

INK = RGBColor(0x2C, 0x25, 0x40)
MUTED = RGBColor(0x6E, 0x66, 0x89)
CORAL = RGBColor(0xF2, 0x60, 0x3F)
CREAM = RGBColor(0xFD, 0xF8, 0xF3)
LILAC = RGBColor(0xF3, 0xEE, 0xFF)
MINT = RGBColor(0xE9, 0xF8, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Inches(13.333), Inches(7.5)

# The speaker script for each slide, so the .pptx is presentable on its own.
NOTES = {
    1: "Open with the number: five forms, five organisations, no one telling you which matter. "
       "Then: 'we built the thing that does that list for you'.",
    2: "The 24-hour freeze window is the hook. A platform report is not a legal record - a bank will not "
       "reverse without a complaint, and a court needs an FIR. Most people find that out too late.",
    3: "Walk the nine stages quickly. The line that matters: 'the agent reads the form it actually has today, "
       "not the one we screenshotted last year'.",
    4: "Two seconds per row. The point is that every row is a real workflow we support today, not a roadmap item.",
    5: "Be honest about the model: it is small, local, and explainable. 'money contributed +2.1 to fraud' is a "
       "sentence a person can check. The agent does the part an LLM cannot - read a live page.",
    6: "The mandatory slide. Four surfaces, each with a 'what it replaces'. Finish on the sandbox toggle: "
       "'a judge can verify this in thirty seconds'.",
    7: "Three decisions only. The stateless one is the engineering story judges like: serverless forgets between "
       "requests, so we made the whole app a pure function of the request.",
    8: "Do not read this slide. Run the demo. The timings on this slide are what you say while you wait.",
    9: "The gap analysis is the product. Most victims do not know what they have not covered.",
    10: "Free for the victim, paid for the organisations. Show the real unit cost - about Rs 25-50 per case - "
        "and say that a Rs 499 plan keeps 90 paise in every rupee.",
    11: "It is already running. Give the URL. Say the measured numbers: 2 seconds to find the portals, 70 seconds "
         "to read the form, 100 seconds to file it.",
    12: "Keep it to three items. Item 2 is the good one: outcome tracking turns a reporting tool into a market-size "
         "argument.",
    13: "One sentence. Mobile-first with authenticated filing, an open channel graph, and two NGO pilots.",
    14: "Only if asked. The point is the core has no dependencies and the evaluation is reproducible.",
    15: "Do not present this. It is here so the form can be filled quickly.",
    16: "Human in the loop twice. The agent will not guess a required field. That is a feature, and it is shown in "
         "the demo.",
    17: "The appendix for the question 'how does a first-time user actually use this'. Each statement has an "
         "action attached to it - that was the design rule.",
}



def parse(md: str):
    """Split the deck into (title, [body lines], [notes]) blocks."""
    slides, cur, in_notes = [], None, False
    for raw in md.splitlines():
        line = raw.rstrip()
        if line.startswith("## "):
            if cur:
                slides.append(cur)
            cur = {"title": line[3:].strip(), "body": [], "notes": []}
            in_notes = False
            continue
        if cur is None:
            continue
        if line.strip().lower() == "speaker notes:":
            in_notes = True
            continue
        if line.startswith("---") and not line.strip("-"):
            in_notes = False
            continue
        if in_notes:
            cur["notes"].append(line)
        else:
            cur["body"].append(line)
    if cur:
        slides.append(cur)
    return slides


def add_text(slide, text, left, top, width, height, size, color, bold=False, align=None, spacing=1.0):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    para = frame.paragraphs[0]
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Inter"
    if align is not None:
        para.alignment = align
    para.line_spacing = spacing
    return box


def add_bullets(slide, items, left, top, width, height, size=15, color=INK):
    """items: list of (text, level) - level 0 is a bullet, 1 is a sub-line."""
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    first = True
    for text, level in items:
        para = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        run = para.add_run()
        run.text = text
        run.font.size = Pt(size if level == 0 else size - 1.5)
        run.font.color.rgb = color if level == 0 else MUTED
        run.font.name = "Inter"
        run.font.bold = level == 0
        para.space_after = Pt(7 if level == 0 else 3)
        para.line_spacing = 1.15
        para.level = level
        marker = "•  " if level == 0 else "–  "
        bullet = para.add_run()
        bullet.text = marker
        bullet.font.size = Pt(size if level == 0 else size - 1.5)
        bullet.font.color.rgb = CORAL if level == 0 else MUTED
        bullet.font.name = "Inter"
        bullet.font.bold = level == 0
    return box


def tint(slide, color):
    shape = slide.shapes.add_shape(1, 0, 0, W, H)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def bar(slide, color, top=0, height=Inches(0.11)):
    shape = slide.shapes.add_shape(1, 0, top, W, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def build():
    md = SOURCE.read_text(encoding="utf-8")
    slides = parse(md)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    for index, block in enumerate(slides, start=1):
        title = block["title"]
        # "Slide 2 - Problem statement" reads badly on a slide; keep it in the
        # markdown for navigation, drop it in the deck.
        display_title = re.sub(r"^Slide \d+\s*[-–]\s*", "", title)
        body = [l for l in block["body"] if l.strip()]
        slide = prs.slides.add_slide(blank)
        tint(slide, CREAM)
        bar(slide, CORAL)

        is_title = title.lower().startswith("slide 1 -")
        is_appendix = title.lower().startswith("appendix")

        if is_title:
            tint(slide, LILAC)
            bar(slide, CORAL, height=Inches(0.16))
            add_text(slide, "RECLAIMYOURNET", Inches(1.0), Inches(2.0), Inches(11), Inches(0.5), 15, CORAL, bold=True)
            add_text(slide, "ReclaimYourNet", Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.3), 54, INK, bold=True)
            add_text(
                slide,
                "Report once. File everywhere. Prove it happened.",
                Inches(1.0), Inches(3.7), Inches(11.3), Inches(0.6), 24, MUTED,
            )
            add_text(
                slide,
                "One description of online harm becomes every report you need to file - on the right portals,\n"
                "with the right evidence, and with proof that you actually filed them.",
                Inches(1.0), Inches(4.5), Inches(11.3), Inches(1.0), 16, INK, spacing=1.3,
            )
            add_text(
                slide,
                "KriticalHacks 2026  ·  AI/ML track  ·  built with TinyFish      |      live: reclaim-your-net.vercel.app",
                Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.4), 13, MUTED,
            )
        else:
            accent = MINT if is_appendix else WHITE
            if is_appendix:
                bar(slide, MINT)
            add_text(slide, display_title, Inches(0.85), Inches(0.45), Inches(11.6), Inches(0.8), 30, INK, bold=True)

            # Split the body into a code block (kept monospace) and bullets.
            items, code = [], []
            in_code = False
            for line in body:
                if line.strip().startswith("```"):
                    in_code = not in_code
                    continue
                if in_code:
                    code.append(line)
                    continue
                if not line.strip():
                    continue
                if line.startswith("#"):
                    continue
                if line.startswith("- "):
                    items.append((line[2:].strip(), 0))
                elif line.startswith("  ") and items:
                    items.append((line.strip(), 1))
                elif line.startswith("|") and set(line.strip()) <= set("|-: "):
                    continue
                elif line.startswith("|"):
                    cells = [c.strip() for c in line.strip().strip("|").split("|")]
                    items.append(("   ".join(c for c in cells if c), 1))
                else:
                    items.append((line.strip(), 0))

            top = Inches(1.45)
            if code:
                box = slide.shapes.add_textbox(Inches(0.85), top, Inches(11.6), Inches(0.2))
                frame = box.text_frame
                frame.word_wrap = True
                for i, line in enumerate(code):
                    para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                    run = para.add_run()
                    run.text = line
                    run.font.size = Pt(11.5)
                    run.font.name = "Consolas"
                    run.font.color.rgb = RGBColor(0x44, 0x3C, 0x5E)
                    para.space_after = Pt(0)
                    para.line_spacing = 1.0
                top = top + Inches(0.09 * max(len(code), 1) + 0.25)

            if items:
                size = 15 if len(items) <= 9 else 13
                add_bullets(slide, items, Inches(0.85), top, Inches(11.6), H - top - Inches(0.6), size=size)

        # speaker notes
        notes = "\n".join(l for l in block["notes"] if l.strip())
        if not notes:
            notes = NOTES.get(index, "")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    prs.save(OUT)
    print(f"wrote {OUT}  ({len(slides)} slides)")


if __name__ == "__main__":
    build()
